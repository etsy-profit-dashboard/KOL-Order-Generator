import io
import base64
import zlib
import re
import unicodedata
from pathlib import Path
from difflib import SequenceMatcher

import requests
import pandas as pd
import streamlit as st

st.set_page_config(page_title="WMS Thai Address Converter", page_icon="📦", layout="wide")

POSTAL_DATA_URL = (
    "https://raw.githubusercontent.com/"
    "thailand-geography-data/thailand-geography-json/main/src/geography.json"
)
FUZZY_THRESHOLD = 0.86
AUTO_ACCEPT_FUZZY_THRESHOLD = 0.93

def clean_text(x):
    if x is None or pd.isna(x):
        return ""
    s = unicodedata.normalize("NFC", str(x))
    # Remove invisible / zero-width characters from WMS / Excel.
    for ch in [
        "\u200b", "\u200c", "\u200d", "\u200e", "\u200f",
        "\u202a", "\u202b", "\u202c", "\u202d", "\u202e",
        "\u2060", "\ufeff",
    ]:
        s = s.replace(ch, "")
    s = s.replace("\xa0", " ").replace("\n", " ").replace("\r", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", s).strip()

def norm_thai(x):
    s = clean_text(x).lower()
    if not s:
        return ""
    replacements = {
        "เเ": "แ", "ํา": "ำ", "ํ": "",
        "": "่", "": "้", "": "๊", "": "๋", "": "์",
    }
    for old, new in replacements.items():
        s = s.replace(old, new)

    for pattern in [
        r"^จังหวัด", r"^กรุงเทพมหานครจังหวัด", r"^อำเภอ", r"^เขต",
        r"^แขวง", r"^ตำบล", r"^ตําบล", r"^จ\.", r"^อ\.", r"^ต\.",
    ]:
        s = re.sub(pattern, "", s)

    aliases = {
        "กรุงเทพฯ": "กรุงเทพมหานคร",
        "กรุงเทพ": "กรุงเทพมหานคร",
        "กทม": "กรุงเทพมหานคร",
        "กทม.": "กรุงเทพมหานคร",
        "กรุงเทพเทพมหานคร": "กรุงเทพมหานคร",
        "จังหวัดกรุงเทพเทพมหานคร": "กรุงเทพมหานคร",
        "อยุธยา": "พระนครศรีอยุธยา",
        "จังหวัดอยุธยา": "พระนครศรีอยุธยา",
        "ลําปาง": "ลำปาง",
        "ลาปาง": "ลำปาง",
    }
    s = aliases.get(s, s)
    return re.sub(r"[\s\-_–—/\\.,;:|()\[\]{}]+", "", s)

def clean_location_name(text):
    if text is None or pd.isna(text):
        return ""
    text = clean_text(text)
    text = re.sub(
        r"^(จังหวัด|จ\.|อำเภอ|อ\.|เขต|แขวง|ตำบล|ต\.|ตําบล)\s*",
        "", text, flags=re.IGNORECASE
    )
    return {
        "กรุงเทพ": "กรุงเทพมหานคร",
        "กทม": "กรุงเทพมหานคร",
        "กทม.": "กรุงเทพมหานคร",
        "กรุงเทพฯ": "กรุงเทพมหานคร",
        "กรุงเทพเทพมหานคร": "กรุงเทพมหานคร",
        "Bangkok": "กรุงเทพมหานคร",
        "BANGKOK": "กรุงเทพมหานคร",
        "bangkok": "กรุงเทพมหานคร",
    }.get(text.strip(), text.strip())

def thai_part(x):
    x = clean_text(x)
    return x.split("/", 1)[0].strip() if "/" in x else x

def english_part(x):
    x = clean_text(x)
    return x.split("/", 1)[1].strip() if "/" in x else ""

def make_bilingual(th, en):
    th, en = clean_text(th), clean_text(en)
    if not th:
        return ""
    return f"{th}/ {en}" if en else th

def clean_mobile_number(val):
    if val is None or pd.isna(val):
        return ""
    text = str(val).strip()
    if "e+" in text.lower():
        try:
            text = str(int(float(text)))
        except Exception:
            pass
    else:
        text = text.split(".")[0]
    text = re.sub(r"\D", "", text)
    if text.startswith("66") and len(text) >= 11:
        text = "0" + text[2:]
    elif len(text) == 9:
        text = "0" + text
    return text

def valid_postal(x):
    if x is None or pd.isna(x):
        return ""
    m = re.search(r"(?<!\d)(\d{5})(?!\d)", clean_text(x))
    return m.group(1) if m else ""

def extract_postal_code(address):
    return valid_postal(address)

def get_wms_postal(row):
    # WMS already contains the receiver postal code in many exports.
    for col in ["Receiver Zipcode", "Receiver Zip", "Postal Code", "Zipcode", "Zip Code"]:
        if col in row.index:
            code = valid_postal(row.get(col, ""))
            if code:
                return code
    return ""

def extract_address_components(address):
    address = clean_text(address)
    result = {"province": "", "district": "", "subdistrict": ""}
    if not address:
        return result

    thai_name = r"([ก-๙เแโใไฤฦะาฯ\-]+)"

    for pattern in [rf"(?:จังหวัด|จ\.)\s*{thai_name}"]:
        matches = re.findall(pattern, address)
        if matches:
            result["province"] = max(matches, key=len)
            break

    for pattern in [
        rf"(?:อำเภอ|อ\.)\s*{thai_name}",
        rf"(?:เขต)\s*{thai_name}",
    ]:
        matches = re.findall(pattern, address)
        if matches:
            result["district"] = max(matches, key=len)
            break

    for pattern in [
        rf"(?:ตำบล|ต\.)\s*{thai_name}",
        rf"(?:ตําบล)\s*{thai_name}",
        rf"(?:แขวง)\s*{thai_name}",
    ]:
        matches = re.findall(pattern, address)
        if matches:
            result["subdistrict"] = max(matches, key=len)
            break
    return result

def similarity(a, b):
    a, b = norm_thai(a), norm_thai(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    return SequenceMatcher(None, a, b).ratio()

def exact_match(items, value):
    key = norm_thai(clean_location_name(value))
    if not key:
        return None
    for item in items:
        if item.get("key") == key:
            return item
    return None

def fuzzy_match(items, candidates):
    best_item, best_score, best_candidate = None, 0.0, ""
    usable = [clean_location_name(x) for x in candidates if clean_location_name(x)]
    if not usable:
        return None, 0.0, ""
    for item in items:
        if not item.get("key"):
            continue
        for candidate in usable:
            score = similarity(item["key"], candidate)
            if score > best_score:
                best_score, best_item, best_candidate = score, item, candidate
    if best_item is not None and best_score >= FUZZY_THRESHOLD:
        return best_item, best_score, best_candidate
    return None, 0.0, ""

@st.cache_data(ttl=86400, show_spinner=False)
def load_postal_data():
    r = requests.get(POSTAL_DATA_URL, timeout=30)
    r.raise_for_status()
    return r.json()

# =========================================================
# EMBEDDED ADDRESS MASTER
# =========================================================
# The address master is stored inside app.py.
# Daily users only need to upload the WMS Excel file.
EMBEDDED_ADDRESS_B64 = (
    "eNqtfd1SW0my7r2eggewg2cwZvd4H36GQPQmfFk2CmuFkBYhtA7huxZDbNzu7r3nnIbGgh6mBUGoRwwTRsbT4uLEeRXe4LzCqfytrFpLdE/MREzH2FatWrWq"
    "srIyv/wy62H67mF6+zA9fJheP0xH+OejJw/3Xz1M9x+mx/MP05uH6Sf8+0/+z7WHaf9hevkw/fgwPX2Yfn64fzc/t9R1rzJ6xv84wR+vsKeB78L/i2+Szy27"
    "Ts/VpMl3D/DiAfYyxR4/+1au87rpOu5V0c10DL7xObYc+gZNl8+t+SZ5W3v668N0/DC9wyYfoLP7b6GnbbebdZ48TA9wLDf4hSd+uPhn7Go777yZ22hKV/4j"
    "+/iqQ3yz/8N7Py7oqr3TdA3fdq3Z6Pkez7BH/+YLbEWTcYGfubbtMt+s697Cl+7juwbYbx9fPYD35h3fZ6NDHzh8mP4iwyvP7ErRcP7FOMG1B16qAXxING/P"
    "/Yz0mjxtMLhP2MS3PZ338972X+l2/PNfQ+/w1vc4YfShJ9jLBHt53cz9h+66nAZ3jS+j8YV5nJ9b9cPiCaxhZ1fpePIOD+YGv3uCouVb/d1/5fzcMy8NRaNF"
    "z97i5A2w0QmOIus4mOcjeXSA0/RRpGGIYzvGOfJfdDE/t+g6cyuu1YQvzdzcYrbb62ave7Z/L7z32M0dveKt22kW7Sf4+TT7P2Lntp2XDVx15/8r2tTbFH88"
    "R8l47v91x7d5IpJFg3uHH3sIzzu/gg4miSf9M47lkGYch/68mcEKr7ss7gU+t4/Lgr24lit1cv8HXKEhSgr3swL93B/wqt7vY9OD+bm6l8uXOcz4Mb78lgbY"
    "9c88gbXjfXEui+xHtgA7bTvP9Bk/riN8BhbnXJblBL/WCzZtSz9c/8X0zAQXe8O1KkV9lhIQmTe6oIaNDrDFQLae73rVr7nfTavubc7voI25L5rsOxD/3Ivq"
    "Ho5JO/Fb+I+oBbiHNdfzesB/1s80ClQp1MUYBNHvFtAq9cxP5J7r2a5UN0lXOA1P8EHSq/s4rI/xV9b9N7aa225uIfk+muj3+PxnnLSJ9r3uen5OdrO2eyKb"
    "k6bzDJ+D2fVtUWi7RcP2qt/1QV4y1Fdp9/7zvLC1246W+Bq314gmm1UTzPE7HNKLwouUV0x+QLl9E00cveMAZ0169zMHW6WP78aFAkEdo7B5JQsa2/FUlFWc"
    "/6qg4fr4w9ei9Se6nx6m35P6L143vRbpSn/08wd8YIR/HsO4/Hc2sz342HDW0NbFg2aT9NtA9RA8w2fKp3irL4Hu4b0O4/67kfK/+r1UNEik/cP/idNK68bq"
    "iU6FN6iYava7Za8O4dvgw0Bauq6N2pEOl6mObdFPMpwENZSHE9FVH0SH+lkEOS/wOOhE+uaADxZQGCekcpbwTMlrIlwf8U0jHN2BDO0zztotn70wwDWQiMLt"
    "+Oezboa9PJFF8o/dwxT6fy203wn2+7mkBrCjrINfrD0MeRHmVnL+zFvUNMfxN/bw86p1TqW5oDoHrYaaOdfLO+gG5WyKT/Op7w98s4eevS16Xlbfwia9wHce"
    "4givRIW+hHdt+Hf16E20raQv/6A/fkF7v5ORg7zNz62DiICg0VOkXA7lKf/b6htHeodOSH9++l/92VDP21uNnj51gkOZotz6B2Gs2wWeA2PUKqT0knWtO99y"
    "F85Ynp47XLc7Ohygo8wvV486oZVS4a77vcGiTQ/+BfYpTD2pxj4+7nvvFNugyCPVRnv9UBeAFdySH7Gf7Jq11RIB8tbaa1UZVpn9CPsT/nrIaswf7w05pW1f"
    "9/gHUinaXSfujaQMZrrPva2AZcO25CUeziAgjSf4uX05GMiK+KwSBjbyLpiaGYtSjY0RMqvhBcfQU9FCI5SUFc4i61uyuPs4u5FR7Zes8Lqs6wWutSvn6RBn"
    "eSLLhcoa181vBr/Ufs26aFAlFvS+kR22oP3/Faj2hvhDH1XTRLq/hM3V2nJNlxrjV6K3uJ9lkKz7vujVvxn5H5Oug6E15l5A4zpbXUHxqx1Hx/b/JpVBlg8Z"
    "cpso6h2wzC9wLcZ0HMKG3M17aMip6YQbC6cDOvA7q42PsWszgWe23ROj61Xsb/EzD3DuSe97+XetVhOm/VKUPwnODa7rEeztbO7f2LmYiLjfySSdiOG95lDq"
    "vdmAn1QzWooarHvjA7fyuUzNwPgVx6LvL0gmO3CA91xWM16M9PMW+6nQobP8JNWh1l2qGYPGCiVZMsFhSQZ6KsbMMo3xFVozV/g2281yvhNMAp0Kskf9Yx0y"
    "Rckh4plFh2jHycftyyYYyQhG0g5tgHoBI9ijg506OkfFMKCOmqAPQJgOcJfu488fyGBEaXov53vLT0jOepBN1Qv4ggYM/72cfyfkPXm7N4fJs/boR/Gw/wKW"
    "qNfFqAje45jGOKZD0rV12o9jbHxlNEzdtdBcbpHTIu4Mrie5M3AgNOhZ3e2wiXKyVv1D3zz2vmP85yv4515BBz/IO00a+44OXpLLE6Q0TsQqJ2Hpi+rwvRc9"
    "MAVajqeZjMOwOfA4wC1R7m8sh+dl2MvQH30NmVGnOEAUP1itH1FfLzi047PKLidhNrEzmU3T0/RPIkPQ0xdOPpZMaEID4PyAPgALyJOB9EH9gVqAx1843EZj"
    "szeMF9E1YMmZbNQLmXjy5Nb8zvQ6LOfvuWO0xwto3FnmNy2bWxbX+IXdO2/kvPT6jDqhw3tfcAY45Qs/EV7B8uNX87jM0pp7xOM0fmvh1YWgBeR9/2g86Qm/"
    "vEA/fMWFDmlR/WH3X6hWv40t3XrhNZsaunqQjtR9pk43M10c6vJOnI9LMnv8uGBxSa9/J9Z2MIDhTAJHlezdWmy7HIiyXWXTrcIqrQSJRJcKVpT0e4qjGYlC"
    "+onsav+OV4VjrfQk+VwxI8h5w8+e+x16borQHJqD5hlBY6glSScexjK1iSqeZAoRQrbQ34mLSHPpHfBn7Y5fCa/3ujmCXtIdwxff4Y6j7l467m6KRsRlvKRf"
    "bvnZkBUl7/i92FH36Jx7zQnfrn0QbnEkZuoPuKRfeveg67ayXmJoqo/yidcLHRU8929EAxCWFA0KJF5HZfqDaYh98nW/qNrZKS6eno1Rl6/U1e+V+2Vj6FCM"
    "UNaqbBS9QfNIUdqpAfnOjdGnJ8ICwBBocyfDIr9qQDP2Kt8GO8GbTk+MnB5Kb4dstoFCW8KFfEFYVCOvPSnZjHBY2LP3AJfHz8G3OGq///+bz16/pWASWkWv"
    "ZpzVCPYicCyrPTHA0ETsmAnjMiBXXgxrdIKcyeD1bFKjDDQNrXwbvK3akwR4nGeUsVbexrOgWd7JBqGtPZE1vBBTGQ4CbwAuupwn+BqnbJioa7A4Gm3U07XY"
    "mObZ/EVs5oGxzmEAdv5gXdWUIO1wShO54VqZ26mRVTsRDXZFBxPJCnTrohYyi3wAXSosi5+Fz+CR383w68KDZ4JxHnKzNbBO3tSSszCyqORolnkiKP5Qln2E"
    "CDxpEnyZopRlL4fhSvFzKpb0cbQ8LKyC5rUnMXyHhiUcTAMxeHdpKcYGzxCEWiMuAA2yw+xnu9GrJZEVK2l0SC3lhPskDVVFEtjrW6GKdDrHJCuq8xdy9irs"
    "KhFyOZT9iwvlT5iitFDiCdESoeMTVuBWvvkOV4Cn3x8Lq1kLAjDUcCKfdSPYREdHlC5OOeigC5ITaFSzcK4uywQPinWALWsK0h7jjvuMG6EO8FEzw21g1bk9"
    "rVDbvkS1A+MaiNmgamLVSwQdHqhC2HC9NECetxrybt6u/rIkJBJEDbABVIYcqCspCO/jqn4YBxg2hZ5bkSq6FCD7gj3mnR1AoZIZIFUInzIMR86Ka7PEXItY"
    "HhA8jC8JIvkXOY3JoE7AlSW32wC7vP2q6IQz5utgv6OCltNlqdnNWQESBFrGnJcUcU4bkmPBv+Ov70QGpnLyjTFG1OuBDVkjn+cWhd0KgZ8mIwN0wI7ByEBr"
    "AzRqYwekyR9iqZygY8KmGIpK3T1dfeNS1fcVRzKC/sVIRl4ruwsnYjku4J4Bk5F28acQyLunY+p7gTcOWKG2O37uV715Rgv5k+BIR6I/NOaFNr1rbWUa+qpZ"
    "43pfzpMRGdf+/KYTbZaE22CbEXKKzaXiJ0bvPQ1NJHDBm7yLrqFCWDaSNXBxjSPTp9bdXiMyLPq40h/pMKw7/hFcnXeC4dIGgKhWs0AQt4WrQWfZCcUXCGbP"
    "5ACiR44f7t9j34C4lidFvr96tmzQUecpLMCVuJVDxCRo4LQncLvvZTw3V4IajwUkWUZAqaCNsI9vGgicfgXbpPDipEeLjWDuy8A58rhkFLUNUhLgOdB2dZbN"
    "IzGBaFwURln0VtByEeaBZhRDqTkv8ZlZrSN0ktdQAHo83d/hjwrK2lgrYutzK76jbggSz2Xm9KFDp7wIM8O3uhoSxWXtRc4rH9WEW8oRrb/+WbawnQFo+YUL"
    "k6Ctr8TCgRbLainSj6TERzQk/7tf0SxtMAm/qtYaxQul/if5h0u0VEnDgOtTK0S0Kt3NZMamgrnLczRnwR+wL9FlRrOSNSA96Nd7jZRgH0dyJebf9xLfs/jv"
    "71xn2y97B+J87MJMmQRibA3wHCNbw45djXQeMtvpDO5V4VZ4VBNsZWWdeAus+EHWN+g7dC8YlI+BYWhVB6ivvGfQfuM9g90o6vSFU+PrUoRsRBBO/qvCHflX"
    "sXyvxFIFf3iHwxiqiDfaaYtjlesNVtNBLAkk4KfXs+jZAMTCj5vpfrixMYHfd9LfgrQ/myHs10YWX8gBogswFM1T9xZO+PFWUOKAgqNBFWDvqOGJURXQinZe"
    "+P3M+GQ32pcX8I1ceyK/5SfF+3aaos4nIp2H8+RxL/JTJNxehr4RuOb37VaOj/0N/8UL1ou8F7YQ6Xz2E4CnRf2gBcokkkPypHPRwuFMvZJ5BfXptrMqBRrR"
    "TliokH1SwyBpIPocC4GDRPYlmDOIw4h3dSuY3wSdql235YJBTOGfO3z7zwiUg9z4Rd6lDTSWH/vsCbSoXxsdvNAZKpA00wuzTfb2J1nU72DaAQFud8veF0EZ"
    "R+h3LRUlL+697MtD9d9iWGRsZJAiEa5d4fhbi5RdfjJK1cG7P5AwpHfw1lVXlxeHgtNhcXopGHGI7U8Zhlh9k+9EuIICUxfyYWhH0ZDK/qIShNBkdS1WWKcK"
    "PiOAvINaq3rM9LgMGHuIVAs5rvxZoGC8bdRLj8Ub/bme96wkEY4pJw+fdHnNRLihxSG7TDAf9OuNnHRDQbL9THwJ2qE8BwM17FddTaN/fSFiIZsB4kbbVatW"
    "SYni2bDMKB6UhkIkSgzW/+8hEgIQCOsCG9xTyy2E9GLjbSABnyGRTzbIaNdxqZ0fQvc0LDwe65kFGjgiNhCIgd5zLsDrB/x0f8D/L/ZCkGLRBlJQ3nlsbpTp"
    "FU8MEb7sYpqIKa0nBkgjAaeVJbN0QNL9Za9JI/0ganxqKGG+RU/oqqpO+0xdA/Mq4JYnglVuNE1j4uSyDPofN12F73euqMuankDKUT2TbbK2DWZTi0HQ8xl+"
    "85r1m2dOqn7gCsa3/Zzm3c4OfyYD399TlFSA5vVGpxBig1XXqitwh4mesGHPNISXhfhdChVYXG2g2pUlF0C1Ve5fiKQw7R9F9sYEMTSJL7sNCiFqOzbRO9ob"
    "QvzyBppr7MrcH/APuf71MmVkGIjqnWxIZeP53dFqOda/XyEOeUV2UojLRfDtJQoeEVu38waJ7txSRuCF9R9HbOxioFdcjH1sc2AM++cQrWa7PqLhqod4QsYp"
    "IhIFonS9CAiZlDQNwCGiaTSQG/ny6saHn4mNQQE0WJTM/hqC0BgCVkgyxMbI8FrAoFhmf4yOTfg9nJkWfaGT7xfDGoddtuw1SZs32aU4/uxhAqOQujqTz/8Q"
    "O02g0DcC288q/fdKeaFmz0WEzs2LLCMM15k6E7DxXCRFXN6s7R49PCr5pnZfG9ZpRT+6f1m/yg5WErqa0wMzCcxJbzcgdJpFCsHGEyygtCmRBEGVeMAYPpJo"
    "etZi3IV+O2Irb0uw3ns5ahiKm0exQQyOv03x1LF8WD1n/VGFME0RcBiKkl3FtVjIc+Sj/4b2hEWZ1oFX1kcs8Sqw1EApiY7Zl24pKYIPmHf4gr8KkV3hVb/h"
    "t92rpl9MaPW1ca28DO4Fy/Of1zQfcVAj1A6+zdNCfhFrVzgp+nqQ3W1X9rsv402KHqhs0x8Mq0RJBZfClkcSSwfYv9Hp/TOBWsqpr7ttlyJ8YzZbC46mJr9H"
    "I6JWLzX4eCqb75hDyDveU+s5c5ifisvqT/KFvBPFWfYFKuIoG5ke1h4ciPtEJqE/z6xyQI894iX4pcnbTEuwOiS0ZwWCDYPuOBN6k9cdO8WjdsDjrPLY4rLk"
    "8prltw7kxRQTuJIIyTqS6TOKCASO2Fict2X0p1tO/TqMCcF2f81GxrVoeyb14T4y6FefpZLPfxu3XPIO1zKj1/tyPPeNa7LURUIXyZgesceyznLKbuBKv6kl"
    "bKtxaFGXLX0sa/1nMfi8SzP3hQu/HYhT4H8glyAN6go27RSZPjXhnWOZ6bHoYS9sO/7ohnnedSImxG6+RmGJBcq9LZTpQgDYudlX1MR39Li3UjY6E0ER61NB"
    "JIos3ynas8bq2iJFqFgosNYg4r4KzZEJDy97g/yl+PylWEEIDtQsFdTyXurMpLJ7t2/QV9q+3cLFTdQqCdFWx2GIUsw1Nk5cUDKiDVglXwpGCAphqauR4mTW"
    "Z3ISQkoFnx1k3Q5QvkEzo38iGR/4MwXSLg3xxNsoT10msSORYfiEUxLgBTMTpi/4stfNvHrIlekaKiWStUHQEWbKETYxP/cye9p5I0gRfc4hvmwdA2wN/gXf"
    "xac9DWa9aOwKpD5SnFq2pnVHAv0vY+4fK6CpUFkug23Rgq0SGsD8XeMSw0Z9KuT3euFHDdH+fJu89bgxr/sR89qo7Zrbyv4lp3YKdsfHc2bP56Bt1Fr7UY9T"
    "tM/WmqUj1QJddGQqBWEgprbuLn8w8uZS93coBjgnfhQS2mR06ysZMjGtFnI/XCBaJUiLjbGQEe/UU6atv4a7PkQzzxOq4BpORqXABqKiCmknDQRFUDew4bJ2"
    "RdhjbOIcdReoQgTHHDBVCC2EPY3zKYB7J4K4j0ByLwMR5Ld8o6BnA89MEMFvcHDJmVDPzXEg9HYmex7KRB5qThNavMtBhenCR9ShwlCH1Kg5EHAOFZjbCiGL"
    "hEK4xqBRN9jpdgEeybeKdi9BwwdkJXQzP83b9JVKHh8rCN4Fro747/elaK33ek3Elr5lIL7sEsAwAsb/a+xq4ldSg3bH7BAKrQm1GuJpeYLdW/bUqhPylDrR"
    "R0xKSFzxrMIdZ+HS0/t78byZZmH3m56JzGKGI7GXNjoyOwcaLZb6McEoE3A757xRIYStuEyJAwTiimvOzHzre0JYTkDIELtnh1q4YAzA3opZfyVeMvfhRdmb"
    "99tOexGkEgb5QQw8himz3WoO0yN5fqxHACQC2Fqj/IG/aUNLy5BeI8ElzXbQJHVyN38mtxtT1FechEOOJScbGPuNksval0x0YjBn6e/XrLVQuJR+mkRU1CNA"
    "N3ad3m1Tgaxylmwg0c8fhHltLRjg274V0yWd1lnJjzynNgcyMhKvhac6jDn6YDMu+yeWmM9KE3yF3ko/eCNt3wjUWBQh4MQQQpQhCaRWZgmPI5ZwnQ8vynu7"
    "SJDDIlPg8ETEM+CSODE7BEdG5udYPL9J8DHrbqfpOrObxSezeUSOaA17f1/K3FR4s2JtfnsuqSxYOaU0Au5NTlUI3tWRwmvjd2R/X4txBq7gC7LNBsa+8hpy"
    "sdFIFdXYBNLZgKgH++Ey5mUeW5YpJkh3eqpX+vjRUZZBWgPjVfYr0zcrZdbOV1ZBghT6tzkV0IzFhL811pymNbS7JlKbtMt3NWyMkqK8vB09VioSoUL+U2od"
    "XhgC+ZGxEV8SzwJOca3moLp6IL4ERwm9p9JtMFmbyd9gZr3Nq2ewnDIcNEOPlEJKuxkI7abDjuEQvjxO2QADnsNjlCeo8e6XLkS6WYejaYNzgD9nNtp/kdI1"
    "w563qpeybwVE9KZ6KyKz2fQDsM/XWa+oH26VLrveqnRVI3wjUh9UQBy2U2il6lf+w5V4u4m6OZPZMMjXmmtWtrqRfcutnmad6lDYpcaUMesWJ+VHxuLJpaKM"
    "XG9yvnAxMvbP5Hqz60BGAiEhr4mooUkXARiD/cbHoJ4+p0ZD6brA0YPhC1mYBCaXXBxOxVEIm1NfG0R3itJXOGUlwMVIzSuqKH5kEym/T0nyto0CMIHht1JY"
    "1lOfyyWAJXfsBw2h4hb21lP765DyLtT9YkN+aHn9wCd17beE3VCrb0Qfw4/1aucsSa+XnY5Z9hHI9AdMlWfnLAs67oyDiDvBhlMSh45uKRdU6UB27sc05ld0"
    "ne7jY6GI3Mr0BaWIsAniQoWCKknjM+GQUcs1F9kDZ5Q/xcbAmmbP8EaWyJZ1ZyC0NpuBmhQbCHPYedN5I8bKoeg1rPaEZMiIz6JkAOYUAKXlOdnGPG8jCTEO"
    "RC6Xmph9CwnsHZ0JPpWG9ImunfcihWQttwVmMeYplk5rxzA3welvi46llwawfk3Bev3t3CRLD7EFdFDJ5i+XXAjTJ5UXomNXlVCaiegPCfcGjootAp31iQji"
    "xYOaYV3ehIclf3qRw4/sT7OEnRushHTtBtl7bCapgKlnf4NN0GSPVPZQkPyB4fiAYsDkY4srROVbFFRYBfWovue5DdMjGEDx+fCz5FtHEVkaq4Wf4EmK0O5E"
    "hKtkvZLKFrpeWOCClfB7QfpOZfxD3OaboPW77pU/8dtxXNPWKtAM8k0+NEGDGaO2H6LPXFjCm/Rzy41WZPcearUoJSC1GS6YCBx0LBNiqak9YHcxNVXEzYTa"
    "6XDbaJbzW4aiEXW9MHthxe10U8Ohb1ZiKhxRBsfRAYta6xg/caN1P752Gt06l0mXDC4IcK3BfLdnrmS51EhYT6k4kq7TTSCkKDgvCTfQzyknekjCzULeSxMj"
    "+8Yr4CoXlLpC/kCkGS/MUY5q8SX5VxpO/USxVJoQxWAjll0VwU7rzSDeWUJfDLVPmjTa1ezgyvoqOo9aZqUCHPhO3nemOiOkbrz0enM7L6WDik9BKaB5Tcoj"
    "sI8woAoJ6B3UbImBa9EmI6oysKoJh0mFjnO1OKU8xyO0/1n1YNLPj2hNJs+UBWktyNFdGuiOpnUT1Iwre2/CrZmnSGqjlKCm1QtWuWpdyPrA7RdngC9DbFgD"
    "YdKEAdA+/V53Ah4TUGdRtMVGxzL/7aTZyjc6S10JvoTiV7bWBtU22gvkERKT5Twxe98L22OgnEyNrzCBOjCnle0xFCcN2R4rjU45tU3DiDCP+2EesaYXpu4H"
    "k/AEEeGvYU97TUaEhBA0QZpG5dlSLuej01Mww78fM8dHEtfmKmGYRNx1e9nrMCDNlWfkuujp4dmXPK5PqH/8lmnXbBmiqVGriv8v5btID5MQwHsNwFC+PboL"
    "uPp5lUOgBSPYG3hZNALIwo1pRogFCZ2RTaAwzL0EexaLTohhn4YQDxWE0Wg1Vdz5T5tPVvoqtPhbrSamloWvS5fosbJItFZRdaSaLfE1lhDh1Ph6L6liz66k"
    "Eo5EG362QCUI807G0GRNK8zpCUBlnNp+1Z/+u1p1SeaUBU0tP5cccOTn6pHPabjynrEF0eY5KZepjfmMaZpR6onmiCs+1cqlC+5NuYJCKVWjEuIvZSGY87Pm"
    "svm5//d//8+s/0Xpaxc4KUNOX3tZKFtvEvJYQVdSbGqNVnWnaYJTyNObodseK08lIhKqVFWBSWaqCVUS593y0j8KhKNBSErd6nIUEtqcyvgwbLczt5I3Yq3w"
    "XsQLtMJzlS1+/0h1hjeMN0O461O5nkDsV6FkUGUBda+sfTsMZYDArewENsVFyRdB19P4IjYlRol9kBFDxD7doFQx6aPRjrgzkbcw42RKy34pZZGqf9W0hOid"
    "uCZ94nZimSbEWIC76ULI8FCZKSFwiKdvJuysvjiTklzzKsQ1JlKj9ZgiGnMbeSCpDrRTidNQpkCHXIaeeNenkmG10OhVh5ls1TL9ZArFKh6JBc3oyy5NITMs"
    "XNC2yS2axz1KTg3M4N6UQ+OdSag+FybssfxVGaoD8kcKyGpCNzxr25zsQ6Na0GFrzkpF1XKvilbtCXk/sPfUxrSKk/BmNDZtzIk6fF8KO+0hItcV5EoDTZcB"
    "V0YePQWZGOSTMhTMmLrU6hNgF2UVfqwJNOYqzErbP4lyWnOJ/EUNpIRCXiqhQNqRoh2gMPzO36raLLNK1bEIUcW6iuc0FzVgHfWigtefFqnMpE7lk3KgOU4V"
    "x4BzR5LEGfK8Sto0tiuzyM/V9CAruRSRgww5KTL5ImP4lCIu6kCQZfDD/NwzzP5yvYhEEMF+BvTri1I/ExbVtwQsNipnP6rrJ5OO5f1CNPhGpBtTuZ8Whou0"
    "X8pVqxctk6vWF94MaCXhFUdwnTK3JwGua5PLYE2I8zjflo0Hhrc6JSLOgalBSBUNStT7aRRoU7rQmezavgkXrW0XXnEsmvGXZ/GiNItvU7+iXFRFMr46UcKw"
    "FjnBhGGtcnIkh6Yt8fAR/32E0uYt7w4VetgqFAFTwPJcwhpN1yjPV6j94CqK46ibiSGPkJ454dOLs4Q57SDSNZMITMnRtWhXz2K5gqPOpfX+FSX5SlgThwSR"
    "rDW3Q24sAb22CitWWeQqrDaKmNgKS5yXIjwgaUeSDn+9FHuGRD1Qb28lrjYRo/WK07QX3FaS5zUyuepIbSsS+lt8fjgbN4sYZ0NhTyBJRYJNYbeUSuLQlglV"
    "cdKVKFfA5GVY9taf6JkrsekHEuNdRkPe+46m4stHgbJ+EP5G3TGSlbZQsCu0qHaerc+cxkoO5B4IFuNDhWreaCSG0i2vBS6EdHJvSMtE/KCwdLOaB3aWsPDW"
    "KrPYkmqgMoFUFDSywk8wvIyZ600X/XIuQtlnRnKzm7VSzvJQyKXQYiVPc0vD5QoSweMI7ViiOLTvTsm2X+tiaS4+7S7FiAplavNGgICvTb7F12KZXHKWOFhc"
    "aRxe+RFluiYQJZ5XJa+SbFHBmaxCM/3NaqYXeVy+wwTv3KzYXVJu1axUs+ikKzWxKyV0fi/6Fdn/kfRemaQr4l2uGkZjRUh0YtENIaPYJDkbM18MvCzVHlei"
    "FVB1ZBb6YM7UIaMfURp+X0lV+OuLNKz1nmNRLUFmOPt4P4RQQyK9hlkv5Zz6u0ZSvc2yJQebQlbKJYAWwiWgXqw2hF9tAZcETL2yRVOW2VsdGTfnozDpDjU4"
    "stMsyqlXGqJiO002+o3sGur0+/m5f+tKFagEW7uK8xMFW9vWXOikhIOt3iWNOXFmXywuy17y/nGrmRTWuRZbDJGJFy4CsAOVWTNRuF7u0nbmEiomslKFgbnU"
    "VKDs78K7uGdni24NaFCeWfAZqZRpMgcLmI/HMxDCPJSnwSUNkD3RRdkvkoQRSiB2b4sKj8Za/gmHFHwcsP6ZRipOkjG20E0K1haJXVL5DDkAWaUeKRt3pipj"
    "RftZRZfFm8mwbG+jF6F/mgNk4Z2XUqpJOHNatPYzfufPaIasu90M09Hi7aAyjhuBg7A/x4TRY8nXwKJ2WuDGFg/XMBDWDec4kKVnaqz0RHx+STfcQVgiopBM"
    "DYk5gVi/7HFNKgMeKxJ1KpFgQKIWdI9NDQ49MclfS8Vu0dbibZ9MdpRyyilR24HzEtmCiaSdG0kjuzDjmsMhuYUNTDNvZ5w3iaYmyBU3PjbFcmmJN7xjtPpG"
    "kxbULdKw1I0C2uQZwYb+ssSXtkYUWotFlN5HTnDZYOxYa1GJGsJ6loshlhjdUoAgVCJQhI3KRSicR3vsPCnA3Kn295My47JPTLXxAH0PpGSFze3ZBL/Wf8ie"
    "BckvJfhh/fpN18WisOLbozuhayIhQqW1z1M6GXEJidVu6IRjA+KaSqxa4old/77kj2qhY8r1etbyns2zdqeIswg4wa/iUjJIKmiwNKUE6BuBJ/fFbGIm9Is8"
    "4ywLO3Thy7DTqSAxUWZWBSR+JzCCpHTFWfqSlh/VoRgJERbgrjQZJ6TC4y0s7TToT+aUhPH93mxXi4ytLi/ywkXmQ/oWnXZHdK1BLxQAOo5lAhWfEYgTORk0"
    "WxPLMHIcMirzxBni3gXM9d/DvryIEv9eOYnMm6PHFNfPEJ9UcZeSQVFA4wXmH2YR46lvfEpdwSW3HepHjMSkPjK0qMW8k1YvPCebUmCExMDXQiJk11MtkaEI"
    "CJdmkrTiiiXT4v66v3tFqD8Ije4Rheq8KRISKGMQRAKNqUJa6IYBkO1SVZixJAMeKsnUFKpV8uaZ2X1HajWAqbLlGknbuKg/HwJHvIvImOg6ru//vNmohEd/"
    "9e4CnaVwhUFaWplMks+a9C0u+q+8q1wUJXqXLY8CHdya6h8TqMe+AzB2BFkGkg7O36lcoIEg9oJWE024toFlhyKX/6aRT0qnRLhFIUIgv5bxXSIw8G0AIVuu"
    "SoGSkSNKk3IbLUB0b6ChwqarJnc4bji5wXEWtp3c4aAfoodnZTGLC0S6B6aSxctGJ/V/yqVxyQuSyrhxdUOjfTYVO1IFdCPq8juzO0j3/+zFoAspNDuwxo0S"
    "felA+UhccDf/1ySbhsD0uZBUF9FeLHrV8NkNW0N5WsJAy0KdSpVkrmWAfOxX9LahwMUJm26lUDrdrBWuAtnkto1o0azhtukM+sURMuF+B720mTVdJ4Ksx3IZ"
    "3jViFqDgXriUToVZhuEqiUbFz9+Y4HpFSdxwgmCTRWHeBkJZQsjT+qVLBJQi3sv8tYBnxJBRBBZV95Z0ZCl+8QVeXcTYuzNU4qx7TNRN4wQ63guP5Nz8e8cG"
    "aYgD1BeQgFLnkNLsNWK2relrn01eDTfCBKWVErwWQlSUrH7B1cGwzAhUTGzPzow0pS+uTTIWIkvLCKzNmJ7yDS0ajeOLWqp4x+hJBh1inalbCbpPBDi44Cz4"
    "Ht7UuZ30mIDy2FIh+aQKhLIz0VxvgepoWX1RcUbTdYd0OEfzZeVx0QSQ2N67E0GbGAIyYSYjpOwsgtbiiwuoRkMvTUi5kvw+czMCFFhbcbPWYub9N7ooSLUu"
    "5WxVpZfxVaRryiUNpRfPRcQGhmxZp1Bnhtu+llzRN9X0xOfNovUIhRM2xTeBucmo5g1hQibm/GUwzpJbXo+EkftHAdFQbrayUnXMc7EoWf3gpc9Q+KDqkgMy"
    "MC7ohgNvTWTp/qRT8I+x80x7dd3eQQAdH2gxU4KosSGCuGkgyd6+rBlaIUd5IvcMR8ee8isXMWwfpWxNq7Wo0bPsyMnP7req4RAM0PKq01At3hn8MC3P+Vta"
    "KrPvSEocUEFL3EiVjUMdfm5Z15NgYC7nsGd+qJoD5/yuMemnpoQQIRWdnWbWmbkRH783SrejXh8VjvKfBFkm8sR/ZK+a7n8G8kRSx4E8T1PK4ddrPQhZyFw6"
    "ggwBQPSjAP1IZvFOAJUkWwTvCYAj0HH9kXfWYuPwtXjDGt3Wn/+H+TW6ACCVfAb8Qk5yIG/dSwlMrgsasOvKK3tXESvqJnX6aCKOsEjfazJUg6s1VtMPJCJi"
    "2jBD8lK2G/NtFikxqfWYcNxVQsV4EViEQAxKNxr25LpqDcIehTr77a0iKUeRXMSA564cjz9LPvRlQmsBXdW1YMeA8z86Ws5Ub53+KNdMC+YRjmOpWYgXnuCh"
    "THfE/zO1Qk5C6hQ7/LD4CzEJv7pExz9eZsQ2v7JPPEz/hI8U9Ezxj9Um+S2NYIb+ZA5FYrp/4aS68owaJlBKMeYu2PiIljPG9KakTmuJy2wx9lSQKy+gCyWf"
    "ApndSnKgJask77mdUgLCrDLxeLGQKRH/G2quqCv9B3M7JdPhghQfiHS/EYxbIwpSuIRy4e/fFXM2w0FAaTCY/S+sD9LaYwHctGWPVl0oeTSw9Y64Bm+5Kspt"
    "KEPqkjj3zPFi5Pt3xdaW91MajV9ZyJk3/nF1U6znmmYFYrFJE7wVMxwKTq5LsmcF1+5AwroT5dq1mkXTqX2DJeHZvsmjSznodqA+fttK3qoQdMtNnHm7U+Wd"
    "hCzDz5K0rFDt3dhn7yXQiqXfyUR7nhSOoDwIcrjbLeGgjKWcHH+/3wd0o11IrRuYiBUxWDsSrEoooWS1jSSqOLHEUG/77uVwb1g4t/gmb1c5I4/dsShTY69a"
    "5FPoF7kqITw+DyVsqXGcMm5wSbWTIYkcz5COhpmULHIclxUBpsiGVhVRwSfP5VRI6HTDXF9m6adY+SwVBE/uiDd+L/dIBPZ5IUBbsLZHZjGA0Zpyf+FujIG4"
    "3F/knejCwHu93aCQTc+wHZEohxSFr+ftdN9/lBDRkez+p1VVjsLlaZDSUiLS/KiZLcyCq6grwjH6KETJWKqWGWntZp1tt+NMODAtq1kqywxuIdarXdcK/PaZ"
    "cutXRcdYUOcIqWGJrkpK1qy7PFlgw5WeKRA6Ear/MBCkwEklJNAWHuWy65wX5wKBLXg9BxKzqjPOXqIFj0SkvwmVeDahWGheAbONH4fZpEaTkFOc5FhfSp7I"
    "RldiN1GnA1MGlqo6Q/fKMVL76syWsV3bdlEW5JnKIxwK5axPrUhElNLcrLphnWuC2Z9jJukEDB1Jgqhe7MpLV2W95e5VVk50q1ofrYNlt7OtcatbSfuW2TBc"
    "hOTK5ktlWFZdbmFerRe1zhx65Q2vlc9H9Wg5Y9KpR/84hUMZsPoo0l/52Tigzerlkq8/IBImBrSfdctSrNkNeqPKYQCM31ZUsP46aqSIgWWHHGB6CjJClppa"
    "/Ic29b5eLsRVJzgzExQ9BGtC+eB9cyL+EdzTRttUuLepRmRZUpSTaylKcsO+qWXwWVDUJa6lziCqIi3oAwbQHfCWlbwXwe6UIKUqCm6x8LaC10ntsGFD6jxK"
    "5C/4yI8osJREjwq7R3dda0jpSmktEE+CEnnt9ECyNWb0/QwJZzoIOfA5TedKqAeYOOm2C3v1reb6QejexUmXFzRjdMplscbYN0kEVGC9iBOrNaUTg1ZagBOG"
    "TTdyKBfWW5bAi8uqjpJhrN01sYfPlBUHthBTTGKbMRQZQqqY47I6ofLv2FzRcIemBqXt4uX02c7M/f7IJcyy65O7mFN2ykCqMypXFwgq6Oi1ibAb+U+WDpn4"
    "T5o7mFVogZHgmB+DNWl0wWbWRVuy5ObPJ5D7nRn3B0b7oaIIl4YPcTKtITWWarn7YoGhqsuxTt3cEh0fyU3Ex1zKACOxG8zytNHasRZFoDb1ckBXl5SC1kvc"
    "rs20hQOTIzexhXiaBChyNR5NQwu3bHDGGZhLv+EquCgAYU8V1CfqQ5WLaNoCmmwBjuVm0MVil07AYw6maqlcOe7cdmayT05EAJKr7ZCfo5HfVMEQP2AaMty6"
    "Bee2DdKsNoPgWS4PVW9KqwC28lZVBewbqfscla6Gxi9dqTHxNe60Qs9OlhbDn8qkn5jbxZgmgReOZDGdgo+R0KhRYlvclEpeoC0odS+ijJyDUIi/kRREk6NK"
    "2ChFI7kPWyNCqy4thnYTRpheAKAFR48CFyYJPCmaKSVSUmQtFFoh0Kxopz7pSESM/dBNdVs1kj/CUKb3KPYEbuX1H9qkBNRrt3JKIlUw5HWtAXzcxUqUvV63"
    "0GIUej3nAUfs4T7OjWZcho1P+gGBBJhn38nfdKoqtUXl2EIFtr+Iuo1rCAPL1vjGw8oSiqZu4kUcsxhxlizolz1n7/pmMR0JeoaHRWMvJvz+lwk8C7+DKpNj"
    "QilvYC0HZ7MJb7UW3FJX84e1ZXLJ47LbwlnZq6pqI+I/SSrcbEfl5VTB/IfbEs2iRqtl+6PdqslPSfVQpdzUvUsa7ucZG0KS0OWTMw3vSus17ZE2EQsAmHx5"
    "xxZL5DrpQ842r+RN7POt3Lyj2xW1FG9sLUXNClWC/FhAu2XX4qvLrJW8b2qZsK28F6blH7js/f8D8m8rJg=="

)

def _load_embedded_address_master():
    raw = zlib.decompress(base64.b64decode(EMBEDDED_ADDRESS_B64))
    return pd.read_csv(io.BytesIO(raw), dtype=object).fillna("")

def load_address_master():
    """Load the built-in address master. No address upload is required."""
    return _load_embedded_address_master()



def prepare_master(df):
    df = df.copy()
    required_columns = ["จังหวัด", "เขต/อำเภอ"]
    if not set(required_columns).issubset(df.columns):
        raise ValueError("ไฟล์ adress.xlsx ต้องมีคอลัมน์ จังหวัด และ เขต/อำเภอ")

    df["จังหวัด"] = df["จังหวัด"].ffill()
    df["province_th"] = df["จังหวัด"].map(thai_part)
    df["province_en"] = df["จังหวัด"].map(english_part)
    df["district_th"] = df["เขต/อำเภอ"].map(thai_part)
    df["district_en"] = df["เขต/อำเภอ"].map(english_part)
    df["pkey"] = df["province_th"].map(norm_thai)
    df["dkey"] = df["district_th"].map(norm_thai)

    province_master = (
        df[["pkey", "province_th", "province_en"]]
        .drop_duplicates(subset=["pkey"], keep="first")
        .set_index("pkey")
    )
    district_master = (
        df[["dkey", "district_th", "district_en"]]
        .drop_duplicates(subset=["dkey"], keep="first")
        .set_index("dkey")
    )
    return df, province_master, district_master

def build_geo_tables(records):
    geo = pd.DataFrame(records)
    columns = [
        "provinceNameTh", "provinceNameEn", "districtNameTh",
        "districtNameEn", "subdistrictNameTh", "subdistrictNameEn",
    ]
    for c in columns:
        if c not in geo.columns:
            geo[c] = ""
        geo[c] = geo[c].fillna("").astype(str).map(clean_text)

    geo["pkey"] = geo["provinceNameTh"].map(norm_thai)
    geo["dkey"] = geo["districtNameTh"].map(norm_thai)
    geo["skey"] = geo["subdistrictNameTh"].map(norm_thai)

    if "postalCode" not in geo.columns:
        geo["postalCode"] = ""
    geo["postalCode"] = geo["postalCode"].fillna("").astype(str).map(valid_postal)
    return geo

def build_provinces(geo):
    out = []
    for (pkey, pth), g in geo.groupby(["pkey", "provinceNameTh"], sort=False):
        if pkey:
            out.append({"key": pkey, "th": pth, "en": g.iloc[0]["provinceNameEn"]})
    return out

def build_districts(gprov):
    out = []
    for (dkey, dth), g in gprov.groupby(["dkey", "districtNameTh"], sort=False):
        if dkey:
            out.append({"key": dkey, "th": dth, "en": g.iloc[0]["districtNameEn"]})
    return out

def build_subdistricts(gprov):
    out = []
    for (skey, sth), g in gprov.groupby(["skey", "subdistrictNameTh"], sort=False):
        if skey:
            out.append({"key": skey, "th": sth, "en": g.iloc[0]["subdistrictNameEn"]})
    return out

def resolve_province(geo, prov_raw, city_raw, area_raw, address):
    provinces = build_provinces(geo)
    parts = extract_address_components(address)
    addr_norm = norm_thai(address)

    p = exact_match(provinces, prov_raw)
    if p is not None:
        return p, "province_exact"

    p = exact_match(provinces, parts["province"])
    if p is not None:
        return p, "address_province"

    embedded = [p for p in provinces if p["key"] and p["key"] in addr_norm]
    if embedded:
        return max(embedded, key=lambda x: len(x["key"])), "address_province_embedded"

    for candidate in [city_raw, area_raw]:
        p = exact_match(provinces, candidate)
        if p is not None:
            return p, "province_from_city_area_exact"

    if prov_raw:
        p, score, _ = fuzzy_match(provinces, [prov_raw])
        if p is not None:
            return p, f"province_fuzzy_auto_{score:.2f}" if score >= AUTO_ACCEPT_FUZZY_THRESHOLD else f"province_fuzzy_check_{score:.2f}"

    if parts["province"]:
        p, score, _ = fuzzy_match(provinces, [parts["province"]])
        if p is not None:
            return p, f"address_province_fuzzy_auto_{score:.2f}" if score >= AUTO_ACCEPT_FUZZY_THRESHOLD else f"address_province_fuzzy_check_{score:.2f}"

    if not prov_raw:
        p, score, _ = fuzzy_match(provinces, [city_raw, area_raw])
        if p is not None:
            return p, f"fallback_province_fuzzy_auto_{score:.2f}" if score >= AUTO_ACCEPT_FUZZY_THRESHOLD else f"fallback_province_fuzzy_check_{score:.2f}"

    return None, "province_not_found"

def resolve_district(gprov, city_raw, area_raw, address):
    districts = build_districts(gprov)
    subdistricts = build_subdistricts(gprov)
    parts = extract_address_components(address)

    province_name = gprov.iloc[0]["provinceNameTh"] if not gprov.empty else ""
    if norm_thai(city_raw) == "เมือง" and province_name:
        city_raw = "เมือง" + province_name

    for candidate, status in [
        (city_raw, "city_exact"),
        (area_raw, "area_exact"),
        (parts["district"], "address_district_exact"),
    ]:
        d = exact_match(districts, candidate)
        if d is not None:
            return d, status

    sub_candidates = [area_raw, city_raw, parts["subdistrict"]]
    for candidate in sub_candidates:
        sub = exact_match(subdistricts, candidate)
        if sub is None:
            continue
        matched = gprov[gprov["skey"] == sub["key"]]
        if matched.empty:
            continue
        dgroups = list(matched.groupby(["dkey", "districtNameTh"], sort=False))
        if len(dgroups) == 1:
            (dkey, dth), g = dgroups[0]
            return {
                "key": dkey, "th": dth, "en": g.iloc[0]["districtNameEn"]
            }, "from_subdistrict_exact"

    for candidate, prefix in [(city_raw, "city"), (area_raw, "area")]:
        if not candidate:
            continue
        d, score, _ = fuzzy_match(districts, [candidate])
        if d is not None:
            addr_norm = norm_thai(address)
            confirmed = d["key"] in addr_norm or similarity(d["th"], parts["district"]) >= 0.95
            if score >= AUTO_ACCEPT_FUZZY_THRESHOLD or confirmed:
                return d, f"{prefix}_fuzzy_confirmed_{score:.2f}"
            return d, f"{prefix}_fuzzy_check_{score:.2f}"

    if parts["district"]:
        d, score, _ = fuzzy_match(districts, [parts["district"]])
        if d is not None:
            return d, f"address_district_fuzzy_auto_{score:.2f}" if score >= AUTO_ACCEPT_FUZZY_THRESHOLD else f"address_district_fuzzy_check_{score:.2f}"

    sub, score, _ = fuzzy_match(subdistricts, sub_candidates)
    if sub is not None:
        matched = gprov[gprov["skey"] == sub["key"]]
        if not matched.empty:
            dgroups = list(matched.groupby(["dkey", "districtNameTh"], sort=False))
            if len(dgroups) == 1:
                (dkey, dth), g = dgroups[0]
                return {
                    "key": dkey, "th": dth, "en": g.iloc[0]["districtNameEn"]
                }, f"from_subdistrict_fuzzy_{score:.2f}"

    return None, "district_not_found"

def resolve_subdistrict(gd, area_raw, city_raw, address):
    subs = build_subdistricts(gd)
    parts = extract_address_components(address)
    candidates = [area_raw, city_raw, parts["subdistrict"]]

    for candidate in candidates:
        sub = exact_match(subs, candidate)
        if sub is not None:
            return sub, "exact"

    sub, score, _ = fuzzy_match(subs, candidates)
    if sub is not None:
        confirmed = sub["key"] in norm_thai(address)
        if score >= AUTO_ACCEPT_FUZZY_THRESHOLD or confirmed:
            return sub, f"fuzzy_confirmed_{score:.2f}"
        return sub, f"fuzzy_check_{score:.2f}"

    return None, "not_found"

def resolve_postal(gd, sub, address, wms_postal=""):
"""
Resolve postal code while validating the WMS zipcode.

Priority:
1) If the address contains an explicit 5-digit zipcode, use it.
   If WMS zipcode is different, treat WMS as incorrect and correct it.
2) If no zipcode is written in the address, validate WMS zipcode
   against the matched subdistrict.
3) If the subdistrict has no single zipcode, validate against the
   district when the district has exactly one zipcode.
4) If there is not enough information to verify the WMS zipcode,
   keep the WMS value but mark it as unverified.
"""
address_postal = extract_postal_code(address)

# ---------------------------------------------------------
# 1. Explicit zipcode in the actual receiver address.
#    This is stronger evidence than the WMS zipcode.
# ---------------------------------------------------------
if address_postal:
    if wms_postal and wms_postal != address_postal:
        return address_postal, "WMS_MISMATCH_ADDRESS_CORRECTED"
    if wms_postal:
        return address_postal, "WMS_VALIDATED_ADDRESS"
    return address_postal, "ADDRESS_EXPLICIT"

# ---------------------------------------------------------
# Build the postal codes belonging to the matched subdistrict.
# ---------------------------------------------------------
sub_vals = []
if sub is not None:
    sub_vals = sorted({
        str(x).strip()
        for x in gd.loc[gd["skey"] == sub["key"], "postalCode"].dropna()
        if re.fullmatch(r"\d{5}", str(x).strip())
    })

# ---------------------------------------------------------
# 2. Validate WMS zipcode against the matched subdistrict.
# ---------------------------------------------------------
if wms_postal:
    if wms_postal in sub_vals:
        return wms_postal, "WMS_VALIDATED_SUBDISTRICT"

    # If the subdistrict has exactly one known zipcode and WMS
    # does not match it, correct the WMS value automatically.
    if len(sub_vals) == 1:
        return sub_vals[0], "WMS_MISMATCH_SUBDISTRICT_CORRECTED"

# ---------------------------------------------------------
# District-level fallback.
# ---------------------------------------------------------
district_vals = sorted({
    str(x).strip()
    for x in gd["postalCode"].dropna()
    if re.fullmatch(r"\d{5}", str(x).strip())
})

# ---------------------------------------------------------
# 3. Validate/correct WMS using a district with one unique zipcode.
# ---------------------------------------------------------
if wms_postal:
    if wms_postal in district_vals:
        return wms_postal, "WMS_VALIDATED_DISTRICT"

    if len(district_vals) == 1:
        return district_vals[0], "WMS_MISMATCH_DISTRICT_CORRECTED"

    # Not enough information to safely replace the WMS value.
    return wms_postal, "WMS_UNVERIFIED"

# ---------------------------------------------------------
# 4. No WMS zipcode: use geographic master data.
# ---------------------------------------------------------
if len(sub_vals) == 1:
    return sub_vals[0], "EXACT_SUBDISTRICT"

if len(district_vals) == 1:
    return district_vals[0], "DISTRICT_UNIQUE"

if len(district_vals) > 1:
    return "", "DISTRICT_HAS_MULTIPLE_POSTAL_CODES"

return "", "NO_POSTAL"

def resolve_row(row, geo, province_master, district_master):
    prov_raw = clean_location_name(row.get("Receipt Province", ""))
    city_raw = clean_location_name(row.get("Receipt City", ""))
    area_raw = clean_location_name(row.get("Receipt Area", ""))
    address = clean_text(row.get("Consignee Addr", ""))
    wms_postal = get_wms_postal(row)

    province, pstatus = resolve_province(geo, prov_raw, city_raw, area_raw, address)

    if province is None:
        postal = wms_postal or extract_postal_code(address)
        return {
            "จังหวัด": "", "เขต/อำเภอ": "", "รหัสไปรษณีย์": postal,
            "สถานะตรวจสอบ": "PROVINCE_NOT_FOUND",
            "จังหวัดสถานะ": pstatus, "อำเภอสถานะ": "",
            "ตำบลสถานะ": "", "รหัสไปรษณีย์สถานะ": "WMS_EXPLICIT" if wms_postal else "ADDRESS_EXPLICIT" if postal else "NO_POSTAL",
        }

    pkey = province["key"]
    gprov = geo[geo["pkey"] == pkey].copy()

    district, dstatus = resolve_district(gprov, city_raw, area_raw, address)

    if district is None:
        province_out = make_bilingual(province.get("th", ""), province.get("en", ""))
        if pkey in province_master.index:
            pm = province_master.loc[pkey]
            province_out = make_bilingual(pm["province_th"], pm["province_en"])
        postal = wms_postal or extract_postal_code(address)
        return {
            "จังหวัด": province_out, "เขต/อำเภอ": "", "รหัสไปรษณีย์": postal,
            "สถานะตรวจสอบ": "DISTRICT_NOT_FOUND",
            "จังหวัดสถานะ": pstatus, "อำเภอสถานะ": dstatus,
            "ตำบลสถานะ": "", "รหัสไปรษณีย์สถานะ": "WMS_EXPLICIT" if wms_postal else "ADDRESS_EXPLICIT" if postal else "NO_POSTAL",
        }

    dkey = district["key"]
    gd = gprov[gprov["dkey"] == dkey].copy()

    sub, sstatus = resolve_subdistrict(gd, area_raw, city_raw, address)
    postal, postal_status = resolve_postal(gd, sub, address, wms_postal)

    province_out = make_bilingual(province["th"], province["en"])
    if pkey in province_master.index:
        pm = province_master.loc[pkey]
        province_out = make_bilingual(pm["province_th"], pm["province_en"])

    district_out = make_bilingual(district["th"], district["en"])
    if dkey in district_master.index:
        dm = district_master.loc[dkey]
        district_out = make_bilingual(dm["district_th"], dm["district_en"])

    unresolved = (
        pstatus.endswith("_check")
        or "_fuzzy_check_" in pstatus
        or dstatus.endswith("_check")
        or "_fuzzy_check_" in dstatus
        or sstatus.startswith("fuzzy_check")
    )

    status = "FUZZY_MATCH_CHECK" if unresolved else (
        "OK" if postal_status in {
"WMS_EXPLICIT",
"WMS_VALIDATED_ADDRESS",
"WMS_VALIDATED_SUBDISTRICT",
"WMS_VALIDATED_DISTRICT",
"WMS_MISMATCH_ADDRESS_CORRECTED",
"WMS_MISMATCH_SUBDISTRICT_CORRECTED",
"WMS_MISMATCH_DISTRICT_CORRECTED",
"ADDRESS_EXPLICIT",
"EXACT_SUBDISTRICT",
"DISTRICT_UNIQUE",
} else postal_status
    )

    return {
        "จังหวัด": province_out,
        "เขต/อำเภอ": district_out,
        "รหัสไปรษณีย์": postal,
        "สถานะตรวจสอบ": status,
        "จังหวัดสถานะ": pstatus,
        "อำเภอสถานะ": dstatus,
        "ตำบลสถานะ": sstatus,
        "รหัสไปรษณีย์สถานะ": postal_status,
    }



# =========================================================
# KOL Excel output
# =========================================================

from copy import copy
from datetime import datetime
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DEFAULT_KOL_TEMPLATE = "ไฟล์ทำ KOL ระบบใหม่ update app(1).xlsx"

def find_kol_template(uploaded=None):
    """Use uploaded template when provided; otherwise use the template next to app.py."""
    if uploaded is not None:
        return uploaded
    path = Path(__file__).parent / DEFAULT_KOL_TEMPLATE
    if path.exists():
        return path
    raise FileNotFoundError(
        f"ไม่พบไฟล์ต้นแบบ {DEFAULT_KOL_TEMPLATE} ในโฟลเดอร์เดียวกับ app.py"
    )

def _clear_values(ws, start_row, end_row, start_col=1, end_col=None):
    end_col = end_col or ws.max_column
    for row in ws.iter_rows(
        min_row=start_row, max_row=end_row,
        min_col=start_col, max_col=end_col
    ):
        for cell in row:
            cell.value = None

def _copy_row_style(ws, source_row, target_row, max_col=None):
    """Copy formatting only so newly added rows look like the template."""
    max_col = max_col or ws.max_column
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)

def _extend_formula_rows(ws, source_row, first_new_row, last_new_row):
    """Copy formula/style row and translate relative references for extra rows."""
    from openpyxl.formula.translate import Translator

    if last_new_row < first_new_row:
        return

    for r in range(first_new_row, last_new_row + 1):
        _copy_row_style(ws, source_row, r)
        for c in range(1, ws.max_column + 1):
            src = ws.cell(source_row, c)
            dst = ws.cell(r, c)
            if isinstance(src.value, str) and src.value.startswith("="):
                try:
                    dst.value = Translator(
                        src.value, origin=src.coordinate
                    ).translate_formula(dst.coordinate)
                except Exception:
                    dst.value = src.value

def write_order_data_to_template(template_source, final_df):
    """
    Create the NEW KOL workbook:
    - Keep the original workbook structure/styles/formulas.
    - Replace only the data area of 'ข้อมูลออเดอร์' with the converted WMS data.
    - Keep all other template sheets and formulas.
    """
    wb = load_workbook(template_source)

    if "ข้อมูลออเดอร์" not in wb.sheetnames:
        raise ValueError("ไฟล์ต้นแบบไม่มีชีท 'ข้อมูลออเดอร์'")
    ws = wb["ข้อมูลออเดอร์"]

    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    missing = [h for h in final_df.columns if h not in headers]
    if missing:
        raise ValueError(
            "คอลัมน์จากไฟล์ WMS ไม่ตรงกับไฟล์ต้นแบบ: " + ", ".join(map(str, missing))
        )

    # Remove old order values but keep the template's formatting.
    if ws.max_row >= 2:
        _clear_values(ws, 2, ws.max_row, 1, ws.max_column)

    data_rows = len(final_df)
    needed_last_row = data_rows + 1

    # Copy the template row style to any rows beyond the original template size.
    style_source = min(max(2, ws.max_row), 2)
    if needed_last_row > ws.max_row:
        _copy_row_style(ws, style_source, ws.max_row + 1, ws.max_column)
        for r in range(ws.max_row + 1, needed_last_row + 1):
            _copy_row_style(ws, style_source, r, ws.max_column)

    # Write values by matching header names, not by hard-coded positions.
    header_to_col = {headers[c - 1]: c for c in range(1, len(headers) + 1)}
    for r_idx, (_, row) in enumerate(final_df.iterrows(), start=2):
        for col_name, value in row.items():
            c = header_to_col[col_name]
            cell = ws.cell(r_idx, c)
            if pd.isna(value):
                value = None
            cell.value = value

    # Make Excel recalculate all template formulas when the file is opened.
    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    return wb

def _patch_new_kol_weight_formulas(wb, data_rows):
    """Make the NEW KOL weight formulas robust for zero-weight SKUs and text quantities."""
    if "ไฟล์อัพโหลด KOL ระบบเก่า" not in wb.sheetnames:
        return
    ws = wb["ไฟล์อัพโหลด KOL ระบบเก่า"]
    last_row = min(ws.max_row, data_rows + 2)
    for r in range(3, last_row + 1):
        # AE = helper line weight. Keep blank when master weight is blank;
        # otherwise safely convert both weight and quantity to numbers.
        ws.cell(r, 31).value = (
            f'=IF(ข้อมูลน้ำหนักสินค้า!C{r-1}="","",'
            f'IFERROR(VALUE(ข้อมูลน้ำหนักสินค้า!C{r-1})*VALUE(J{r}),0))'
        )
        # AA = package weight by address + 0.2
        ws.cell(r, 27).value = f'=SUMIF(W:W,W{r},AE:AE)+0.2'

def _build_weight_lookup(wb):
    """SKU -> single-item weight from the template's 'ข้อมูลน้ำหนักสินค้า'."""
    if "ข้อมูลน้ำหนักสินค้า" not in wb.sheetnames:
        return {}

    ws = wb["ข้อมูลน้ำหนักสินค้า"]
    # Template source table is D (sku) and L (single-item weight).
    lookup = {}
    for r in range(2, ws.max_row + 1):
        sku = clean_text(ws.cell(r, 4).value)
        weight = ws.cell(r, 12).value
        if not sku:
            continue
        try:
            if weight in (None, ""):
                continue
            lookup[sku] = float(weight)
        except Exception:
            continue
    return lookup

def _build_old_kol_values(final_df, template_wb):
    """
    Evaluate the OLD KOL upload sheet's formulas as values only.
    The logic intentionally follows the current template formulas:
    - order data starts at row 2
    - old KOL data starts at row 3
    - district/zipcode fields follow the existing template references
    """
    weight_lookup = _build_weight_lookup(template_wb)
    now_value = datetime.now()

    rows = []
    # Weight helper (AE) is SKU weight * allocated quantity.
    helper_weights = []

    for _, row in final_df.iterrows():
        erp = clean_text(row.get("ERP No.", ""))
        product_code = clean_text(row.get("Product Code", ""))
        barcode = clean_text(row.get("Product Barcode", ""))
        qty_raw = row.get("Allocated Qty", "")
        try:
            qty = float(qty_raw) if clean_text(qty_raw) else 0
            if qty.is_integer():
                qty = int(qty)
        except Exception:
            qty = 0

        raw_address = row.get("Consignee Addr", "")
        address = "" if pd.isna(raw_address) else str(raw_address)
        product_weight = weight_lookup.get(product_code, 0.0)
        line_weight = product_weight * qty
        helper_weights.append((address, line_weight))

        comment = clean_text(row.get("Comment by seller", ""))
        if comment == "Express delivery":
            special_mark = "NORMAL"
        elif comment == "Lalamove":
            special_mark = "B2B_OFFLINE"
        else:
            special_mark = ""

        # Follow the existing template formula references exactly.
        out = [
            "EastlyncOffline" if erp else "",
            "OFFLINE" if erp else "",
            "SALE" if erp else "",
            erp,
            special_mark,
            "EXPRESS" if erp else "",
            product_code,
            barcode,
            product_code,
            qty if erp else "",
            0 if erp else "",
            0 if erp else "",
            now_value if erp else "",
            now_value if erp else "",
            clean_text(row.get("Receive Name", "")),
            clean_mobile_number(row.get("Receiver Mobile Number", "")) if erp else "",
            clean_mobile_number(row.get("Receiver Mobile Number", "")) if erp else "",
            "TH" if erp else "",
            clean_text(row.get("จังหวัด", "")),
            clean_text(row.get("เขต/อำเภอ", "")),
            clean_text(row.get("รหัสไปรษณีย์", "")),
            clean_text(row.get("รหัสไปรษณีย์", "")),
            address,
            "COD" if erp else "",
            None,
            None,
            None,   # filled after grouping by address
            None,
            None,
            None,
            line_weight if erp else "",
        ]
        rows.append(out)

    # Existing template AA formula: SUMIF(W:W,W3,AE:AE)+0.2
    address_totals = {}
    for address, weight in helper_weights:
        address_totals[address] = address_totals.get(address, 0.0) + weight

    for i, row in enumerate(rows):
        address = row[22]  # W
        if row[3]:
            row[26] = address_totals.get(address, 0.0) + 0.2
        else:
            row[26] = ""

    return rows

def create_old_kol_file(template_wb, final_df):
    """Create OLD KOL upload workbook with values only and A:AD only.

    AE is an internal weight helper in the template and is intentionally
    excluded from the standalone OLD KOL upload file.
    """
    src_ws = template_wb["ไฟล์อัพโหลด KOL ระบบเก่า"]
    upload_max_col = 30  # A:AD only; AE is not exported.

    from openpyxl import Workbook
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "ไฟล์อัพโหลด KOL ระบบเก่า"

    for c in range(1, upload_max_col + 1):
        letter = openpyxl.utils.get_column_letter(c)
        dim = src_ws.column_dimensions[letter]
        out_ws.column_dimensions[letter].width = dim.width
        out_ws.column_dimensions[letter].hidden = False

    for r in (1, 2):
        for c in range(1, upload_max_col + 1):
            s = src_ws.cell(r, c)
            d = out_ws.cell(r, c)
            d.value = s.value
            if s.has_style:
                d._style = copy(s._style)
            d.number_format = s.number_format
            d.alignment = copy(s.alignment)
            d.protection = copy(s.protection)

    values = _build_old_kol_values(final_df, template_wb)

    for i, row_values in enumerate(values, start=3):
        for c, value in enumerate(row_values[], start=1):
            out_ws.cell(i, c).value = value

    style_row = 3
    for r in range(3, 3 + len(values)):
        for c in range(1, upload_max_col + 1):
            s = src_ws.cell(style_row, c)
            d = out_ws.cell(r, c)
            if s.has_style:
                d._style = copy(s._style)
            d.number_format = s.number_format
            d.alignment = copy(s.alignment)
            d.protection = copy(s.protection)

    out_ws.freeze_panes = src_ws.freeze_panes
    out_ws.auto_filter.ref = f"A1:AD{max(2, 2 + len(values))}"

    try:
        out_wb.calculation.calcMode = "manual"
    except Exception:
        pass

    return out_wb

def workbook_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

# =========================================================
# UI
# =========================================================

st.title("📦 WMS → KOL Excel Generator")
st.caption("แปลงข้อมูล WMS → ใส่ลงไฟล์ต้นแบบ KOL ระบบใหม่ → สร้างไฟล์ KOL ระบบเก่าแบบไม่มีสูตร Excel")

with st.sidebar:
    st.header("ตั้งค่า")

    st.success("Address master: ฝังอยู่ในโปรแกรมแล้ว")
    st.info(
        "ใช้งานโดยอัปโหลดเฉพาะไฟล์ WMS เท่านั้น\n\n"
        "ไฟล์ Address master ถูกเก็บไว้ภายใน app.py "
        "ไม่ต้องอัปโหลดทุกครั้ง"
    )

    st.info(
        "ไฟล์ที่ 1: ใช้ไฟล์ต้นแบบ KOL ระบบใหม่ โดยวางข้อมูล WMS ที่แปลงแล้วลงในชีท "
        "'ข้อมูลออเดอร์' และคงสูตร/รูปแบบของไฟล์ต้นแบบไว้\n\n"
        "ไฟล์ที่ 2: สร้างชีท 'ไฟล์อัพโหลด KOL ระบบเก่า' แยกออกมา โดยแปลงสูตรเป็นค่าข้อมูลจริง"
    )

uploaded = st.file_uploader(
    "1) อัปโหลดไฟล์ Excel จาก WMS",
    type=["xlsx", "xls"]
)

if uploaded is None:
    st.markdown("""
### วิธีใช้งาน
1. อัปโหลดไฟล์ **Outbound Detail Export จาก WMS**
2. ระบบตรวจสอบ `Receipt Province`, `Receipt City`, `Receipt Area`
3. ตรวจสอบ `Consignee Addr`
4. ตรวจสอบ `Receiver Zipcode`
5. สร้าง `จังหวัด`, `เขต/อำเภอ`, `รหัสไปรษณีย์`
6. นำข้อมูลที่ได้ไปใส่ในชีท **ข้อมูลออเดอร์** ของไฟล์ต้นแบบโดยอัตโนมัติ
7. สร้างไฟล์ **KOL ระบบใหม่** 1 ไฟล์ และ **KOL ระบบเก่าแบบไม่มีสูตร** อีก 1 ไฟล์
""")
    st.stop()

try:
    raw = pd.read_excel(uploaded, dtype=str).fillna("")
except Exception as e:
    st.error(f"อ่านไฟล์ไม่สำเร็จ: {e}")
    st.stop()

required = ["Receipt Province", "Receipt City", "Receipt Area", "Consignee Addr"]
missing = [c for c in required if c not in raw.columns]
if missing:
    st.error("ไม่พบคอลัมน์ที่จำเป็น: " + ", ".join(missing))
    st.stop()

if raw.empty:
    st.warning("⚠️ ไฟล์ Excel ไม่มีข้อมูล Order")
    st.stop()

try:
    master, province_master, district_master = prepare_master(load_address_master())
except Exception as e:
    st.error(f"อ่าน address master ไม่สำเร็จ: {e}")
    st.stop()

try:
    geo = build_geo_tables(load_postal_data())
except Exception as e:
    st.error(f"ไม่สามารถโหลดฐานข้อมูลรหัสไปรษณีย์ได้ในขณะนี้ ({e})")
    st.stop()

with st.spinner("กำลังประมวลผลข้อมูล WMS..."):
    results = [
        resolve_row(row, geo, province_master, district_master)
        for _, row in raw.iterrows()
    ]
    result_df = pd.DataFrame(results)

insert_at = list(raw.columns).index("Receipt Area") + 1
final = pd.concat(
    [
        raw.iloc[:, :insert_at].copy(),
        result_df[["จังหวัด", "เขต/อำเภอ", "รหัสไปรษณีย์"]],
        raw.iloc[:, insert_at:].copy(),
    ],
    axis=1,
)

if "Receiver Mobile Number" in final.columns:
    final["Receiver Mobile Number"] = final["Receiver Mobile Number"].apply(clean_mobile_number)

final["รหัสไปรษณีย์"] = (
    final["รหัสไปรษณีย์"].fillna("").astype(str)
    .str.extract(r"(\d{5})", expand=False).fillna("")
)

st.success(f"ประมวลผลแล้ว {len(final):,} แถว")
c1, c2, c3, c4 = st.columns(4)
c1.metric("ทั้งหมด", f"{len(final):,}")
c2.metric("สำเร็จ", f"{(result_df['สถานะตรวจสอบ'] == 'OK').sum():,}")
c3.metric("ต้องตรวจสอบ", f"{(result_df['สถานะตรวจสอบ'] != 'OK').sum():,}")
c4.metric("รหัสไปรษณีย์ว่าง", f"{result_df['รหัสไปรษณีย์'].fillna('').astype(str).str.strip().eq('').sum():,}")

empty_district = result_df["เขต/อำเภอ"].fillna("").astype(str).str.strip().eq("")
empty_province = result_df["จังหวัด"].fillna("").astype(str).str.strip().eq("")
empty_postal = result_df["รหัสไปรษณีย์"].fillna("").astype(str).str.strip().eq("")
issues = result_df[result_df["สถานะตรวจสอบ"] != "OK"]

issue_indices = sorted(
    set(issues.index)
    | set(result_df.index[empty_district])
    | set(result_df.index[empty_province])
    | set(result_df.index[empty_postal])
)

if issue_indices:
    with st.expander("⚠️ รายการที่ควรตรวจสอบก่อนใช้งาน", expanded=True):
        preview_cols = [
            "จังหวัด", "เขต/อำเภอ", "รหัสไปรษณีย์",
            "สถานะตรวจสอบ", "จังหวัดสถานะ", "อำเภอสถานะ",
            "ตำบลสถานะ", "รหัสไปรษณีย์สถานะ",
        ]
        left_cols = ["Receipt Province", "Receipt City", "Receipt Area", "Consignee Addr"]
        if "Receiver Zipcode" in raw.columns:
            left_cols.append("Receiver Zipcode")

        issue_preview = pd.concat(
            [
                raw.loc[issue_indices, left_cols].reset_index(drop=True),
                result_df.loc[issue_indices, preview_cols].reset_index(drop=True),
            ],
            axis=1,
        )
        st.dataframe(issue_preview, use_container_width=True, hide_index=True)

st.subheader("ตัวอย่างข้อมูลที่แปลงแล้ว")
show_cols = [
    "ERP No.", "Receipt Province", "Receipt City", "Receipt Area",
    "จังหวัด", "เขต/อำเภอ", "รหัสไปรษณีย์",
    "Product Code", "Product Barcode", "Allocated Qty",
]
st.dataframe(
    final[[c for c in show_cols if c in final.columns]].head(100),
    use_container_width=True,
    hide_index=True,
)

try:
    template_source = find_kol_template(None)

    with st.spinner("กำลังสร้างไฟล์ KOL ระบบใหม่และ KOL ระบบเก่า..."):
        # Load once for both outputs.
        template_wb = load_workbook(template_source)

        # File 1: full template + converted WMS data in 'ข้อมูลออเดอร์'.
        new_kol_wb = write_order_data_to_template(template_source, final)
        _patch_new_kol_weight_formulas(new_kol_wb, len(final))
        new_kol_bytes = workbook_bytes(new_kol_wb)

        # File 2: standalone old KOL sheet, values only.
        old_kol_wb = create_old_kol_file(template_wb, final)
        old_kol_bytes = workbook_bytes(old_kol_wb)

    st.success("สร้างไฟล์เรียบร้อยแล้ว")

    col1, col2 = st.columns(2)

    with col1:
        st.download_button(
            "⬇️ ดาวน์โหลดไฟล์ 1 — KOL ระบบใหม่",
            data=new_kol_bytes,
            file_name="KOL_New_WMS_Updated.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )
        st.caption(
            "ใช้ไฟล์ต้นแบบเดิม + ใส่ข้อมูล WMS ลงชีท 'ข้อมูลออเดอร์' "
            "โดยคงสูตร Excel และรูปแบบของต้นแบบ"
        )

    with col2:
        st.download_button(
            "⬇️ ดาวน์โหลดไฟล์ 2 — KOL ระบบเก่า (ไม่มีสูตร)",
            data=old_kol_bytes,
            file_name="KOL_Old_Upload_No_Formulas.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="secondary",
            use_container_width=True,
        )
        st.caption(
            "แยกเฉพาะชีท 'ไฟล์อัพโหลด KOL ระบบเก่า' และแปลงสูตรเป็นค่าข้อมูลจริง "
            "ไม่มีสูตร Excel"
        )

except Exception as e:
    st.error(f"สร้างไฟล์ KOL ไม่สำเร็จ: {e}")
    st.exception(e)
